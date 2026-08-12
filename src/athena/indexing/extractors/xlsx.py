"""
Microsoft Excel (.xlsx) document extractor.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from athena.indexing.extractors.base import BaseExtractor
from athena.indexing.models import (
    ExtractedDocument,
    ExtractedPage,
)


class XLSXExtractor(BaseExtractor):
    """Extract text from Excel workbooks."""

    @property
    def supported_extensions(
        self,
    ) -> tuple[str, ...]:
        """
        Return supported file extensions.
        """

        return (".xlsx",)

    def extract(
        self,
        document: Path,
    ) -> ExtractedDocument:
        """
        Extract workbook content as searchable text.

        Args:
            document:
                Path to the XLSX document.

        Returns:
            ExtractedDocument.
        """

        workbook = load_workbook(
            filename=document,
            read_only=True,
            data_only=True,
        )

        sections: list[str] = []

        for sheet in workbook.worksheets:
            sections.append(
                f"Sheet: {sheet.title}"
            )

            for row in sheet.iter_rows(
                values_only=True,
            ):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:
                    sections.append(
                        " | ".join(values)
                    )

        text = "\n".join(
            sections,
        )

        return ExtractedDocument(
            document_id=document.name,
            path=document,
            title=document.stem,
            text=text,
            pages=(
                ExtractedPage(
                    page_number=1,
                    text=text,
                ),
            ),
            page_count=1,
        )
